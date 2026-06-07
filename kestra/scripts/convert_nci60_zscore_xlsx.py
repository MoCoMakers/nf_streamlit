"""
Convert the DTP NCI-60 z-score workbook to a CSV ready for csv_to_datawarehouse.py.

Source : data/DTP_NCI60_ZSCORE/output/DTP_NCI60_ZSCORE.xlsx  (sheet 'all')
Output : data/DTP_NCI60_ZSCORE/output/DTP_NCI60_ZSCORE.csv
Also   : data/DTP_NCI60_ZSCORE/output/DTP_NCI60_ZSCORE_column_map.csv  (original -> sql header)

The sheet has a CellMiner metadata preamble (rows 1-8); the real header is row 9 and
data start at row 10. 68 columns: 6 compound-metadata cols, 60 cell-line z-score cols
(PANEL:CellName), 2 experiment-count cols. Values are written verbatim (raw load = TEXT).
"""
import csv
import re
from pathlib import Path

import openpyxl

SRC = Path("data/DTP_NCI60_ZSCORE/output/DTP_NCI60_ZSCORE.xlsx")
OUT = Path("data/DTP_NCI60_ZSCORE/output/DTP_NCI60_ZSCORE.csv")
MAP = Path("data/DTP_NCI60_ZSCORE/output/DTP_NCI60_ZSCORE_column_map.csv")
HEADER_ROW = 9

# Explicit names for the 6 metadata + 2 trailing columns; cell-line cols are sanitized.
EXPLICIT = {
    "NSC # b": "nsc",
    "Drug name": "drug_name",
    "FDA status": "fda_status",
    "Mechanism of action c": "mechanism_of_action",
    "PubChem SID": "pubchem_sid",
    "SMILES d": "smiles",
    "Total experiments e": "total_experiments",
    "Total after quality control f": "total_after_qc",
}


def sanitize(name: str) -> str:
    """'LC:A549/ATCC' -> 'lc_a549_atcc'."""
    s = re.sub(r"[^0-9a-zA-Z]+", "_", name.strip().lower())
    return s.strip("_")


def main() -> None:
    wb = openpyxl.load_workbook(SRC, read_only=True)
    ws = wb["all"]

    raw_header = next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))
    raw_header = [str(h).strip() if h is not None else "" for h in raw_header]

    sql_header, seen, mapping = [], {}, []
    for orig in raw_header:
        col = EXPLICIT.get(orig, sanitize(orig))
        if col in seen:  # guard against collisions
            seen[col] += 1
            col = f"{col}_{seen[col]}"
        else:
            seen[col] = 0
        sql_header.append(col)
        mapping.append((orig, col))

    n_written = 0
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(sql_header)
        for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
            if all(c is None for c in row):
                continue
            w.writerow(["" if c is None else c for c in row])
            n_written += 1

    with open(MAP, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["original_header", "sql_column"])
        w.writerows(mapping)

    print(f"Wrote {n_written} data rows, {len(sql_header)} columns -> {OUT}")
    print(f"Column map -> {MAP}")


if __name__ == "__main__":
    main()
